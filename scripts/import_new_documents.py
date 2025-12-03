#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
增量导入新文档脚本
用于将data目录中新添加的CSV文档数据导入数据库
"""

import sys
import os

# 添加项目根目录到Python路径
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import csv
from datetime import datetime

from app import create_app, db
from app.models import Organization, Category, Document

# 尝试导入openpyxl用于读取Excel文件
try:
    import openpyxl
    EXCEL_SUPPORT = True
except ImportError:
    EXCEL_SUPPORT = False


def import_new_documents_from_csv(csv_file_path, app, org_name):
    """通用增量文档导入函数"""
    with app.app_context():
        org = Organization.query.filter_by(name=org_name).first()
        if not org:
            # 如果组织不存在，创建它（仅适用于FDA Guidance和APIC）
            if org_name in ['FDA Guidance', 'APIC']:
                org = Organization(name=org_name)
                db.session.add(org)
                db.session.commit()
                print(f"创建了{org_name}组织")
            else:
                print(f"错误: 未找到{org_name}组织")
                return 0, 0
        
        # 读取价格列表（如果存在）
        price_dict = {}
        price_list_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', 'price_list.xlsx')
        if EXCEL_SUPPORT and os.path.exists(price_list_path):
            try:
                workbook = openpyxl.load_workbook(price_list_path)
                worksheet = workbook.active
                # 跳过标题行，从第二行开始读取
                for row in worksheet.iter_rows(min_row=2, values_only=True):
                    if row and len(row) >= 2 and row[0] and row[1] is not None:
                        title = str(row[0]).strip()
                        try:
                            price = float(row[1])
                            price_dict[title] = price
                        except ValueError:
                            # 如果价格不是有效数字，跳过该行
                            continue
            except Exception as e:
                print(f"读取价格列表文件时出错: {e}")
        
        # 检查CSV文件是否存在
        if not os.path.exists(csv_file_path):
            print(f"警告: {org_name} CSV文件不存在 {csv_file_path}")
            return 0, 0
        
        # 获取数据库中已有的所有该组织文档标题
        existing_docs = Document.query.filter_by(org_id=org.id).all()
        existing_titles = set(doc.title for doc in existing_docs)
        
        # 读取CSV文件
        with open(csv_file_path, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            imported_count = 0
            skipped_count = 0
            
            for row in reader:
                title = row['title'].strip()
                
                # 检查文档是否已存在
                if title in existing_titles:
                    skipped_count += 1
                    continue
                
                # 查找或创建分类（可选）
                category = None
                if row.get('category') and row['category'].strip():
                    category = Category.query.filter_by(name=row['category'].strip(), org_id=org.id).first()
                    # 对于FDA Guidance和APIC，如果分类不存在则创建
                    if not category and org_name in ['FDA Guidance', 'APIC']:
                        category = Category(name=row['category'].strip(), org_id=org.id)
                        db.session.add(category)
                        db.session.flush()  # 获取新分类的ID但不提交事务
                        print(f"添加了新的{org_name}分类: {row['category'].strip()}")
                    # 对于其他组织，如果分类不存在则创建
                    elif not category:
                        category = Category(name=row['category'].strip(), org_id=org.id)
                        db.session.add(category)
                        db.session.flush()  # 获取新分类的ID但不提交事务
                        print(f"添加了新的{org_name}分类: {row['category'].strip()}")
                
                # 解析日期
                publish_date = None
                if row.get('publish_date') and row['publish_date'].strip():
                    try:
                        publish_date = datetime.strptime(row['publish_date'].strip(), '%Y-%m-%d').date()
                    except ValueError:
                        pass  # 忽略日期格式错误
                
                # 获取价格信息
                price = price_dict.get(title, 0)  # 默认价格为0
                
                # 创建文档对象（完全统一的字段处理）
                doc = Document(
                    title=title,
                    chinese_title=row.get('chinese_title', '').strip() or None,
                    summary=row.get('summary', '').strip() or None,
                    chinese_summary=row.get('chinese_summary', '').strip() or None,
                    org_id=org.id,
                    category_id=category.id if category else None,
                    cover_url=row.get('cover_url', '').strip() or None,
                    publish_date=publish_date,
                    source_url=row.get('source_url', '').strip() or None,
                    original_file_url=row.get('original_file_url', '').strip() or None,
                    translation_file_url=row.get('translation_file_url', '').strip() or None,
                    price=price
                )
                
                db.session.add(doc)
                existing_titles.add(title)  # 添加到已存在集合中防止重复
                imported_count += 1
                
                # 每100条记录提交一次
                if (imported_count + skipped_count) % 100 == 0:
                    db.session.commit()
                    print(f"已处理{org_name}文档 {imported_count + skipped_count} 条...")
            
            # 提交剩余记录
            db.session.commit()
            print(f"{org_name}文档增量导入完成: 新增 {imported_count} 条记录，跳过 {skipped_count} 条已存在记录")
            return imported_count, skipped_count


def import_new_pda_documents(app):
    """导入新的PDA文档"""
    csv_file_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', 'pda_documents.csv')
    return import_new_documents_from_csv(csv_file_path, app, 'PDA')


def import_new_who_documents(app):
    """导入新的WHO文档"""
    csv_file_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', 'who_documents.csv')
    return import_new_documents_from_csv(csv_file_path, app, 'WHO')


def import_new_ispe_documents(app):
    """导入新的ISPE文档"""
    csv_file_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', 'ispe_documents.csv')
    return import_new_documents_from_csv(csv_file_path, app, 'ISPE')


def import_new_fda_guidance_documents(app):
    """导入新的FDA Guidance文档"""
    csv_file_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', 'fda_guidance_documents.csv')
    return import_new_documents_from_csv(csv_file_path, app, 'FDA Guidance')

def import_new_apic_documents(app):
    """导入新的APIC文档"""
    csv_file_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', 'apic_documents.csv')
    return import_new_documents_from_csv(csv_file_path, app, 'APIC')


def import_all_new_documents():
    """导入所有新增文档"""
    # 创建应用实例
    app = create_app()
    
    total_imported = 0
    total_skipped = 0
    
    print("开始增量导入PDA文档...")
    imported, skipped = import_new_pda_documents(app)
    total_imported += imported
    total_skipped += skipped
    
    print("开始增量导入WHO文档...")
    imported, skipped = import_new_who_documents(app)
    total_imported += imported
    total_skipped += skipped
    
    print("开始增量导入ISPE文档...")
    imported, skipped = import_new_ispe_documents(app)
    total_imported += imported
    total_skipped += skipped
    
    print("开始增量导入FDA Guidance文档...")
    imported, skipped = import_new_fda_guidance_documents(app)
    total_imported += imported
    total_skipped += skipped
    
    print("开始增量导入APIC文档...")
    imported, skipped = import_new_apic_documents(app)
    total_imported += imported
    total_skipped += skipped
    
    print(f"所有新增文档导入完成: 总共新增 {total_imported} 条记录，跳过 {total_skipped} 条已存在记录")


if __name__ == '__main__':
    import_all_new_documents()
